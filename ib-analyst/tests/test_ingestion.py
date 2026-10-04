"""PDF and Excel ingestion, including all the ways a file can be bad."""

import io

import pytest
from openpyxl import Workbook
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from ib_analyst.config import settings
from ib_analyst.errors import (CorruptFileError, EmptyDocumentError, FileTooLargeError, PasswordProtectedError,
                               UnsupportedFileError)
from ib_analyst.ingestion.excel_processor import relative_formula
from ib_analyst.ingestion.loader import ingest_files, read_file


def _pdf_with_text(text="Hello world. Facility amount PKR 100 million.", pages=1) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for _ in range(pages):
        c.drawString(72, 720, text)
        c.showPage()
    c.save()
    return buf.getvalue()


def _blank_pdf() -> bytes:
    w = PdfWriter()
    w.add_blank_page(width=595, height=842)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _encrypted_pdf(password: str) -> bytes:
    w = PdfWriter(clone_from=io.BytesIO(_pdf_with_text()))
    w.encrypt(user_password=password, owner_password="owner")
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


# ---------------- PDF ----------------
def test_pdf_pages_headings_clauses_tables(sample_docs):
    fa = next(d for d in sample_docs if d.filename.startswith("Facility"))
    p2 = fa.page(2)
    assert "22. EVENTS OF DEFAULT" in p2.headings
    assert {"18.1", "22.4", "24.1"} <= set(p2.clause_refs)
    im = next(d for d in sample_docs if d.filename.startswith("Information"))
    assert len(im.pages) == 3 and im.page(2).tables  # sources & uses tables
    assert any("Total project cost" in " ".join(r) for t in im.page(2).tables for r in t)


def test_corrupt_pdf():
    with pytest.raises(CorruptFileError):
        read_file("broken.pdf", b"%PDF-1.4 this is not really a pdf \x00\x01\x02")


def test_password_protected_pdf():
    with pytest.raises(PasswordProtectedError):
        read_file("locked.pdf", _encrypted_pdf("secret"))


def test_pdf_with_empty_password_opens():
    doc = read_file("restricted.pdf", _encrypted_pdf(""))
    assert "Facility amount" in doc.pages[0].text


def test_scanned_or_blank_pdf():
    with pytest.raises(EmptyDocumentError, match="scanned"):
        read_file("scan.pdf", _blank_pdf())


def test_page_limit(monkeypatch):
    monkeypatch.setattr(settings, "max_pdf_pages", 2)
    doc = read_file("long.pdf", _pdf_with_text(pages=5))
    assert len(doc.pages) == 2 and "only the first 2" in doc.warnings[0]


# ---------------- Generic ----------------
def test_unsupported_type_and_empty_file():
    with pytest.raises(UnsupportedFileError):
        read_file("notes.docx", b"PK\x03\x04data")
    with pytest.raises(UnsupportedFileError):
        read_file("empty.pdf", b"")


def test_too_large(monkeypatch):
    monkeypatch.setattr(settings, "max_file_mb", 1)
    with pytest.raises(FileTooLargeError):
        read_file("huge.pdf", b"0" * (2 * 1024 * 1024))


def test_one_bad_file_does_not_stop_others(sample_paths):
    p = sample_paths["Term_Sheet_Thar_Sun.pdf"]
    res = ingest_files([("bad.pdf", b"garbage"), (p.name, p.read_bytes())])
    assert len(res.documents) == 1 and res.errors[0][0] == "bad.pdf"
    assert "damaged" in res.errors[0][1]


# ---------------- Excel ----------------
def test_workbook_structure(sample_docs):
    wb = next(d for d in sample_docs if d.kind == "excel").workbook
    roles = {s.name: s.role for s in wb.sheets}
    assert roles == {"Inputs": "input", "Operations": "calculation", "Debt Schedule": "debt_schedule",
                     "Financial Statements": "financial_statements", "Summary": "output"}
    ops = wb.sheet("Operations")
    assert ops.period_header[0] == ("C", "2027") and len(ops.period_header) == 12
    assert "Inputs" in ops.referenced_sheets
    ebitda = next(r for r in ops.rows if r.label == "EBITDA")
    assert ebitda.kind == "formula" and ebitda.cells[0].formula == "=C8-C12"
    assert round(ebitda.cells[0].value, 2) == 6466.08  # cached value preserved alongside formula
    inputs = wb.sheet("Inputs")
    assert next(r for r in inputs.rows if r.label.startswith("Installed")).kind == "input"


def test_planted_model_errors_detected(sample_docs):
    wb = next(d for d in sample_docs if d.kind == "excel").workbook
    found = {(i.category, i.location) for i in wb.issues}
    assert ("Hard-coded value in formula row", "Operations!H10") in found
    assert ("Inconsistent formula in row", "Operations!K11") in found
    assert ("Constant embedded in formula", "Operations!K11") in found


def test_relative_formula():
    assert relative_formula("=D10-D11", 12, 5) == relative_formula("=E10-E11", 12, 6)
    assert relative_formula("=$C$4*D5", 6, 4) == "=R4C3*R[-1]C[0]"


def test_error_values_and_external_links():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "Year", 2025, 2026, 2027
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "Revenue", 100, "#REF!", "=[Other.xlsx]Sheet1!A1"
    buf = io.BytesIO()
    wb.save(buf)
    doc = read_file("errs.xlsx", buf.getvalue())
    cats = {i.category for i in doc.workbook.issues}
    assert "Excel error values" in cats and "External link" in cats


def test_formulas_without_saved_values_warns():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "Year", 2025, 2026, 2027
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "Revenue", 100, "=B2*1.1", "=C2*1.1"
    buf = io.BytesIO()
    wb.save(buf)
    doc = read_file("nocache.xlsx", buf.getvalue())
    assert not doc.workbook.has_cached_values and "save it" in doc.warnings[0]


def test_legacy_encrypted_corrupt_and_empty_excel():
    with pytest.raises(UnsupportedFileError, match="Save As"):
        read_file("old.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 100)
    with pytest.raises(PasswordProtectedError):
        read_file("locked.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 100)
    with pytest.raises(CorruptFileError):
        read_file("broken.xlsx", b"PK\x03\x04 not a real zip")
    buf = io.BytesIO()
    wb = Workbook()
    wb.active["A1"] = "Just a heading"
    wb.save(buf)
    with pytest.raises(EmptyDocumentError):
        read_file("blank.xlsx", buf.getvalue())
